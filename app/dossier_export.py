"""Génération du dossier de financement au format Word (.docx)."""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.shared import Cm, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import charts_export as charts
from . import endettement as endet_mod
from .schemas import ExportDossierInput, TypeProjet
from .simulation import simuler
from .utils import clean_result

POLICE = "Calibri"
PRIMARY_COLOR = RGBColor(0x1D, 0x6F, 0x5C)
ALERTE_COLOR = RGBColor(0xB0, 0x2A, 0x2A)
GRIS_COLOR = RGBColor(0x99, 0x99, 0x99)
GRIS_LIBELLE = RGBColor(0x5A, 0x5A, 0x5A)
TEXTE_FONCE = RGBColor(0x22, 0x2A, 0x27)
ZEBRA_HEX = "F2F7F5"
BORDURE_HEX = "E2E2E2"
BLANC = RGBColor(0xFF, 0xFF, 0xFF)

# Page A4 paysage : plus de largeur pour les mises en page tableau + graphique.
LARGEUR_PAGE_CM = 29.7
HAUTEUR_PAGE_CM = 21.0
MARGE_CM = 1.8
LARGEUR_CONTENU_CM = LARGEUR_PAGE_CM - 2 * MARGE_CM

LABELS_TYPE_PROJET = {
    "location_longue_duree": "Location longue durée",
    "location_courte_duree": "Location courte durée (type Airbnb)",
    "achat_revente": "Achat-revente",
}
LABELS_STRUCTURE = {
    "personne_physique": "Personne physique",
    "sci_ir": "SCI à l'IR",
    "sci_is": "SCI à l'IS",
}
LABELS_DIFFERE = {
    "partiel": "Différé partiel (intérêts seuls payés)",
    "total": "Différé total (intérêts capitalisés)",
}


def _eur(v: float) -> str:
    return f"{v:,.0f} €".replace(",", " ")


def _pct(v: float, digits: int = 1) -> str:
    return f"{v * 100:.{digits}f} %"


def _label_type_projet(v: str) -> str:
    return LABELS_TYPE_PROJET.get(v, v.replace("_", " "))


def _label_structure(v: str) -> str:
    return LABELS_STRUCTURE.get(v, v.replace("_", " "))


# ---------------------------------------------------------------------------
# Aides bas niveau (mise en forme Word via oxml)
# ---------------------------------------------------------------------------

def _add_bottom_border(paragraph, color: str = "1D6F5C", size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr_el)
    run._r.append(sep)
    run._r.append(end)


def _ajouter_titre(doc: Document, texte: str, niveau: int = 1):
    h = doc.add_heading(texte, level=niveau)
    for run in h.runs:
        run.font.color.rgb = PRIMARY_COLOR
        run.font.name = POLICE
        run.font.bold = True
        run.font.size = Pt(18)
    _add_bottom_border(h, size=10)
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(10)
    return h


def _set_cell_background(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _bordure_bas_cellule(cell, color: str = BORDURE_HEX, size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    bordures = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:color"), color)
    bordures.append(bottom)
    tc_pr.append(bordures)


def _cell_marges(cell, haut: int = 90, bas: int = 90, gauche: int = 120, droite: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    marges = OxmlElement("w:tcMar")
    for cote, valeur in (("top", haut), ("bottom", bas), ("left", gauche), ("right", droite)):
        el = OxmlElement(f"w:{cote}")
        el.set(qn("w:w"), str(valeur))
        el.set(qn("w:type"), "dxa")
        marges.append(el)
    tc_pr.append(marges)


def _ajouter_table_kv(container, lignes: list[tuple[str, str]]):
    """Tableau clé/valeur épuré : pas de grille Word générique, un simple
    filet clair sous chaque ligne, un léger zébrage et une valeur en évidence."""
    table = container.add_table(rows=0, cols=2)
    table.style = "Normal Table"
    table.autofit = True
    for i, (cle, valeur) in enumerate(lignes):
        row = table.add_row()
        cell_cle, cell_valeur = row.cells

        p_cle = cell_cle.paragraphs[0]
        run_cle = p_cle.add_run(cle)
        run_cle.font.name = POLICE
        run_cle.font.size = Pt(10.5)
        run_cle.font.color.rgb = GRIS_LIBELLE

        p_valeur = cell_valeur.paragraphs[0]
        p_valeur.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_valeur = p_valeur.add_run(valeur)
        run_valeur.font.name = POLICE
        run_valeur.font.size = Pt(11)
        run_valeur.font.bold = True
        run_valeur.font.color.rgb = TEXTE_FONCE

        for cell in (cell_cle, cell_valeur):
            _cell_marges(cell)
            _bordure_bas_cellule(cell)
            if i % 2 == 1:
                _set_cell_background(cell, ZEBRA_HEX)
    return table


def _supprimer_bordures(table) -> None:
    tbl_pr = table._tbl.tblPr
    bordures = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{cote}")
        el.set(qn("w:val"), "nil")
        bordures.append(el)
    tbl_pr.append(bordures)


def _ajouter_table_et_graphique(
    doc: Document,
    lignes: list[tuple[str, str]],
    image_bytes: bytes | None,
    largeur_table_cm: float = 13.2,
    largeur_image_cm: float = 11.5,
) -> None:
    """Tableau de chiffres à gauche, graphique correspondant à droite (mise en
    page côte à côte, comme les dossiers de financement présentés en réunion)."""
    if not image_bytes:
        _ajouter_table_kv(doc, lignes)
        return

    conteneur = doc.add_table(rows=1, cols=2)
    conteneur.autofit = False
    conteneur.alignment = WD_TABLE_ALIGNMENT.CENTER
    _supprimer_bordures(conteneur)
    conteneur.columns[0].width = Cm(largeur_table_cm)
    conteneur.columns[1].width = Cm(largeur_image_cm)
    cell_table, cell_image = conteneur.rows[0].cells
    cell_table.width = Cm(largeur_table_cm)
    cell_image.width = Cm(largeur_image_cm)

    _ajouter_table_kv(cell_table, lignes)

    p_image = cell_image.paragraphs[0]
    p_image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_image.add_run()
    run.add_picture(io.BytesIO(image_bytes), width=Cm(largeur_image_cm - 0.5))


def _centrer_verticalement(section) -> None:
    """Centre le contenu entre les marges haut/bas de la page (comme une page
    de titre) — pertinent ici car chaque section ne contient qu'un chapitre."""
    sect_pr = section._sectPr
    valign = OxmlElement("w:vAlign")
    valign.set(qn("w:val"), "center")
    sect_pr.append(valign)


def _mettre_en_page(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(LARGEUR_PAGE_CM * 10)
    section.page_height = Mm(HAUTEUR_PAGE_CM * 10)
    section.left_margin = Cm(MARGE_CM)
    section.right_margin = Cm(MARGE_CM)
    section.top_margin = Cm(MARGE_CM)
    section.bottom_margin = Cm(MARGE_CM)
    _centrer_verticalement(section)


def _configurer_page(doc: Document) -> None:
    _mettre_en_page(doc.sections[0])


def _nouvelle_page_chapitre(doc: Document):
    """Un chapitre par page : nouvelle section démarrant sur une nouvelle
    page, contenu centré verticalement, en-tête/pied hérités de la page de garde."""
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    _mettre_en_page(section)
    # add_section() copie different_first_page_header_footer=True de la page de
    # garde ; sans correction, chaque chapitre (qui est aussi la "première
    # page" de sa propre section) afficherait l'en-tête/pied vide du premier.
    section.different_first_page_header_footer = False
    return section


def _configurer_entete_pied(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.text = "Page "
    _add_field(footer_p, "PAGE")
    footer_p.add_run(" / ")
    _add_field(footer_p, "NUMPAGES")
    for run in footer_p.runs:
        run.font.size = Pt(8)
        run.font.name = POLICE
        run.font.color.rgb = GRIS_COLOR

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.text = "Dossier de financement immobilier"
    for run in header_p.runs:
        run.font.size = Pt(8)
        run.font.name = POLICE
        run.font.color.rgb = GRIS_COLOR


def _mensualite_et_loyers(inp, resultat: dict, is_achat_revente: bool) -> tuple[float, float]:
    if is_achat_revente:
        ar = resultat["achat_revente"]
        return ar["frais_portage_interets"] / inp.duree_portage_mois, 0.0
    return resultat.get("mensualite_credit_hors_assurance", 0.0), resultat["annees"][0]["loyers_bruts"] / 12


def _meilleur_regime(annee1: dict) -> str:
    cfs = annee1["cashflow_apres_impot"]
    return max(cfs, key=cfs.get)


def _cout_apport_emprunt(resultat: dict, is_achat_revente: bool) -> tuple[float, float, float]:
    source = resultat["achat_revente"] if is_achat_revente else resultat
    return (
        source.get("cout_total_acquisition", 0.0),
        source.get("apport_reel", 0.0),
        source.get("montant_emprunte", 0.0),
    )


# ---------------------------------------------------------------------------
# Page de garde et sommaire
# ---------------------------------------------------------------------------

def _ajouter_table_kv_centree(doc: Document, lignes: list[tuple[str, str]], largeur_cm: float = 13.0):
    table = _ajouter_table_kv(doc, lignes)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    l1, l2 = largeur_cm * 0.42, largeur_cm * 0.58
    table.columns[0].width = Cm(l1)
    table.columns[1].width = Cm(l2)
    for row in table.rows:
        row.cells[0].width = Cm(l1)
        row.cells[1].width = Cm(l2)
    return table


def _ajouter_bandeau_titre(doc: Document, titre_texte: str, sous_titre_texte: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    _supprimer_bordures(table)
    table.columns[0].width = Cm(LARGEUR_CONTENU_CM)
    cell = table.rows[0].cells[0]
    cell.width = Cm(LARGEUR_CONTENU_CM)
    _set_cell_background(cell, "1D6F5C")
    _cell_marges(cell, haut=550, bas=550, gauche=200, droite=200)

    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = p1.add_run(titre_texte)
    run1.font.name = POLICE
    run1.font.size = Pt(30)
    run1.font.bold = True
    run1.font.color.rgb = BLANC

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(8)
    run2 = p2.add_run(sous_titre_texte)
    run2.font.name = POLICE
    run2.font.size = Pt(14)
    run2.italic = True
    run2.font.color.rgb = BLANC


def _ajouter_page_de_garde(doc: Document, payload: ExportDossierInput, inp) -> None:
    # Le contenu est centré verticalement sur la page (voir _centrer_verticalement),
    # inutile d'empiler des paragraphes vides pour pousser le bloc vers le bas.
    _ajouter_bandeau_titre(
        doc, "Dossier de financement immobilier", "Analyse de financement et de rentabilité"
    )

    doc.add_paragraph()

    lignes = [("Type de projet", _label_type_projet(inp.type_projet.value))]
    if payload.adresse_bien:
        lignes.append(("Bien concerné", payload.adresse_bien))
    if payload.nom_emprunteur:
        lignes.append(("Emprunteur", payload.nom_emprunteur))
    lignes.append(("Date de génération", datetime.now().strftime("%d/%m/%Y")))
    _ajouter_table_kv_centree(doc, lignes)

    doc.add_paragraph()

    note = doc.add_paragraph(
        "Document généré automatiquement — estimation pédagogique, à faire valider par un "
        "professionnel avant toute décision."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.runs[0].italic = True
    note.runs[0].font.name = POLICE
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = GRIS_COLOR


def _ajouter_sommaire(doc: Document, titres: list[str]) -> None:
    _ajouter_titre(doc, "Sommaire")
    for i, titre in enumerate(titres, start=1):
        p = doc.add_paragraph(f"{i}. {titre}")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.runs[0].font.size = Pt(13)


# ---------------------------------------------------------------------------
# Sections
# ---------------------------------------------------------------------------

def _section_synthese(doc, payload, inp, resultat, is_achat_revente, meilleur_regime, annee1):
    cout_total, apport, montant_emprunte = _cout_apport_emprunt(resultat, is_achat_revente)
    lignes = [
        ("Type de projet", _label_type_projet(inp.type_projet.value)),
        ("Structure juridique", _label_structure(inp.structure_juridique.value)),
        ("Coût total de l'opération", _eur(cout_total)),
        ("Apport personnel", _eur(apport)),
        ("Montant emprunté", _eur(montant_emprunte))
        if montant_emprunte > 0
        else ("Financement", "100 % fonds propres (sans crédit)"),
    ]
    if is_achat_revente:
        ar = resultat["achat_revente"]
        lignes += [
            ("Durée de portage", f"{inp.duree_portage_mois} mois"),
            ("Marge nette prévisionnelle", _eur(ar["marge_nette"])),
            ("Rentabilité de l'opération", _pct(ar["rentabilite_operation_pct"])),
        ]
        image = charts.chart_barres(
            [("Marge brute avant impôt", ar["marge_brute_avant_impot"]), ("Marge nette", ar["marge_nette"])],
            "Marge de l'opération",
        )
    else:
        lignes += [
            ("Régime fiscal le plus favorable (année 1)", meilleur_regime),
            ("Cash-flow net mensuel (ce régime)", _eur(annee1["cashflow_apres_impot"][meilleur_regime] / 12)),
            ("Rendement brut", _pct(resultat.get("rendement_brut", 0))),
            ("Rendement net de charges", _pct(resultat.get("rendement_net_charges", 0))),
        ]
        image = charts.chart_barres(
            [
                ("Rendement brut", resultat.get("rendement_brut", 0)),
                ("Rendement net de charges", resultat.get("rendement_net_charges", 0)),
            ],
            "Rendement du projet",
            formatter=lambda v: _pct(v),
        )
    _ajouter_table_et_graphique(doc, lignes, image)

    if payload.profil is not None:
        mensualite_projet, loyers_mensuels = _mensualite_et_loyers(inp, resultat, is_achat_revente)
        r = endet_mod.calculer_taux_endettement(
            payload.profil.revenus_nets_mensuels_foyer,
            payload.profil.autres_revenus_mensuels,
            payload.profil.mensualites_credits_existants,
            mensualite_projet,
            loyers_mensuels,
        )
        p = doc.add_paragraph()
        run = p.add_run(
            f"Taux d'endettement estimé : {_pct(r.taux_endettement, 1)} — "
            + ("dépasse le seuil HCSF de 35 %." if r.depasse_seuil else "sous le seuil HCSF de 35 %.")
        )
        run.bold = True
        run.font.color.rgb = ALERTE_COLOR if r.depasse_seuil else PRIMARY_COLOR


def _section_presentation(doc, payload, inp, resultat, is_achat_revente):
    lignes = []
    if payload.adresse_bien:
        lignes.append(("Adresse du bien", payload.adresse_bien))
    lignes += [
        ("Type de bien", inp.type_bien.value.capitalize()),
        ("Surface", f"{inp.surface_m2:.0f} m²"),
    ]
    if not is_achat_revente:
        lignes.append(("Bien neuf / VEFA", "Oui" if inp.bien_neuf else "Non"))
    lignes += [
        ("Prix d'achat", _eur(inp.prix_achat)),
        ("Frais de notaire", _eur(inp.frais_notaire)),
        ("Montant des travaux", _eur(inp.montant_travaux)),
    ]
    if inp.montant_mobilier:
        lignes.append(("Montant du mobilier", _eur(inp.montant_mobilier)))
    cout_total, _apport, _emprunt = _cout_apport_emprunt(resultat, is_achat_revente)
    lignes.append(("Coût total de l'opération", _eur(cout_total)))
    if not is_achat_revente:
        lignes.append(
            ("Régime locatif", "Location meublée" if inp.regime_location.value == "meublee" else "Location nue")
        )
        if inp.type_projet == TypeProjet.location_courte_duree:
            lignes.append(("Meublé de tourisme classé", "Oui" if inp.meuble_tourisme_classe else "Non"))

    composition = [
        ("Prix d'achat", inp.prix_achat),
        ("Frais de notaire", inp.frais_notaire),
        ("Travaux", inp.montant_travaux),
    ]
    if inp.montant_mobilier:
        composition.append(("Mobilier", inp.montant_mobilier))
    image = charts.chart_donut(
        [v for _, v in composition], [l for l, _ in composition], "Composition du coût d'acquisition"
    )
    _ajouter_table_et_graphique(doc, lignes, image)


def _section_profil(doc, payload):
    p = payload.profil
    lignes = []
    if payload.nom_emprunteur:
        lignes.append(("Emprunteur", payload.nom_emprunteur))
    lignes += [
        ("Revenus nets mensuels du foyer", _eur(p.revenus_nets_mensuels_foyer)),
        ("Autres revenus mensuels", _eur(p.autres_revenus_mensuels)),
        ("Mensualités de crédits existants", _eur(p.mensualites_credits_existants)),
        ("Total des revenus mensuels retenus", _eur(p.revenus_nets_mensuels_foyer + p.autres_revenus_mensuels)),
    ]
    image = charts.chart_barres(
        [
            ("Revenus mensuels retenus", p.revenus_nets_mensuels_foyer + p.autres_revenus_mensuels),
            ("Mensualités de crédits existants", p.mensualites_credits_existants),
        ],
        "Profil de l'emprunteur",
    )
    _ajouter_table_et_graphique(doc, lignes, image)


def _lignes_credit(inp, resultat, montant_emprunte, is_achat_revente) -> list[tuple[str, str]]:
    lignes = [
        ("Montant emprunté", _eur(montant_emprunte)),
        ("Taux du crédit", _pct(inp.taux_credit_annuel, 2)),
        ("Durée du crédit", f"{inp.duree_credit_annees} ans"),
    ]
    if is_achat_revente:
        lignes.append(("Durée de portage retenue pour les intérêts", f"{inp.duree_portage_mois} mois"))
        return lignes
    lignes.append(("Taux d'assurance emprunteur", _pct(inp.taux_assurance_emprunteur, 2)))
    if inp.differe_type.value != "aucun":
        libelle = LABELS_DIFFERE.get(inp.differe_type.value, inp.differe_type.value)
        lignes.append(("Différé de crédit", f"{libelle} — {inp.differe_duree_mois} mois"))
        lignes.append(
            ("Mensualité 1ère année (hors assurance)", _eur(resultat.get("mensualite_annee1_hors_assurance", 0)) + "/mois")
        )
        lignes.append(
            ("Mensualité en régime de croisière (hors assurance)", _eur(resultat.get("mensualite_credit_hors_assurance", 0)) + "/mois")
        )
    else:
        lignes.append(
            ("Mensualité du crédit (hors assurance)", _eur(resultat.get("mensualite_credit_hors_assurance", 0)) + "/mois")
        )
    return lignes


def _section_financement(doc, inp, resultat, is_achat_revente):
    cout_total, apport, montant_emprunte = _cout_apport_emprunt(resultat, is_achat_revente)
    lignes = [
        ("Coût total d'acquisition", _eur(cout_total)),
        ("Apport personnel", _eur(apport)),
    ]
    if montant_emprunte > 0:
        lignes += _lignes_credit(inp, resultat, montant_emprunte, is_achat_revente)
    else:
        lignes.append(("Mode de financement", "100 % fonds propres (sans crédit)"))
    image = charts.chart_donut([apport, montant_emprunte], ["Apport personnel", "Montant emprunté"], "Plan de financement", total_label="Coût total")
    _ajouter_table_et_graphique(doc, lignes, image)


def _section_charges(doc, inp, annee1, is_meublee, is_lcd):
    loyers_bruts_an1 = annee1["loyers_bruts"]
    frais_gestion = loyers_bruts_an1 * inp.frais_gestion_pct_loyers

    if is_lcd:
        nuitees_louees = 365 * inp.taux_occupation_pct
        _ajouter_table_kv(
            doc,
            [
                ("Prix moyen par nuitée", _eur(inp.prix_nuitee)),
                ("Taux d'occupation annuel retenu", _pct(inp.taux_occupation_pct)),
                ("Nuitées louées par an (estimation)", f"{nuitees_louees:.0f} nuits"),
                ("Recettes locatives brutes (année 1)", _eur(loyers_bruts_an1)),
            ],
        )
        doc.add_paragraph()
    else:
        _ajouter_table_kv(
            doc,
            [
                ("Loyer mensuel hors charges appliqué", _eur(inp.loyer_mensuel_hors_charges)),
                ("Vacance locative retenue", _pct(inp.vacance_locative_pct)),
                ("Recettes locatives brutes (année 1)", _eur(loyers_bruts_an1)),
            ],
        )
        doc.add_paragraph()

    charges_items: list[tuple[str, float]] = [
        ("Copropriété", inp.charges_copropriete_annuelles),
        ("Taxe foncière", inp.taxe_fonciere_annuelle),
        ("Assurance PNO", inp.assurance_pno_annuelle),
        ("Entretien", inp.entretien_annuel),
        ("Gestion locative", frais_gestion),
    ]
    lignes = [
        ("Charges de copropriété", _eur(inp.charges_copropriete_annuelles)),
        ("Taxe foncière", _eur(inp.taxe_fonciere_annuelle)),
        ("Assurance PNO", _eur(inp.assurance_pno_annuelle)),
        ("Entretien annuel", _eur(inp.entretien_annuel)),
        ("Frais de gestion locative", _eur(frais_gestion) + f" ({_pct(inp.frais_gestion_pct_loyers)})"),
    ]
    if is_lcd:
        frais_plateforme = loyers_bruts_an1 * inp.frais_plateforme_pct
        charges_items.append(("Commission plateforme", frais_plateforme))
        charges_items.append(("Ménage", inp.frais_menage_annuel))
        lignes.append(
            ("Commission plateforme (Airbnb/Booking...)", _eur(frais_plateforme) + f" ({_pct(inp.frais_plateforme_pct)})")
        )
        lignes.append(("Frais de ménage annuels", _eur(inp.frais_menage_annuel)))
    if is_meublee:
        charges_items.append(("Comptable", inp.frais_comptable_annuel))
        lignes.append(("Frais comptable annuel", _eur(inp.frais_comptable_annuel)))
    lignes.append(("Total des charges hors crédit (année 1)", _eur(annee1["charges_hors_credit"])))

    # Le graphique inclut en plus le crédit (contrairement au tableau, exprimé
    # hors crédit) pour donner une vision complète de la rentabilité du projet.
    charges_graphique = charges_items + [("Crédit (annuité)", annee1["mensualite_totale_credit"])]
    charges_graphique_non_nulles = [(l, v) for l, v in charges_graphique if v > 0]
    image = charts.chart_recettes_charges(
        loyers_bruts_an1, charges_graphique_non_nulles, "Recettes et charges annuelles (crédit inclus)"
    )
    _ajouter_table_et_graphique(doc, lignes, image, largeur_table_cm=11.5, largeur_image_cm=13.2)


def _section_achat_revente_detail(doc, inp, resultat):
    ar = resultat["achat_revente"]
    lignes = [
        ("Durée de portage", f"{inp.duree_portage_mois} mois"),
        ("Intérêts de portage", _eur(ar["frais_portage_interets"])),
        ("Taxe foncière de portage", _eur(ar["frais_portage_taxe_fonciere"])),
        ("Assurance de portage", _eur(ar["frais_portage_assurance"])),
        ("Total des frais de portage", _eur(ar["frais_portage_total"])),
        ("Prix de revente visé", _eur(ar["prix_revente"])),
        ("Frais d'agence à la revente", _eur(ar["frais_agence_revente"])),
        ("Produit net de la vente", _eur(ar["produit_net_vente"])),
        ("Marge brute avant impôt", _eur(ar["marge_brute_avant_impot"])),
        ("Régime fiscal applicable", ar["regime_fiscal"]),
        ("Impôt total", _eur(ar["impot_total"])),
        ("Marge nette", _eur(ar["marge_nette"])),
        ("Rentabilité de l'opération", _pct(ar["rentabilite_operation_pct"])),
    ]
    if ar.get("tri_annualise") is not None:
        lignes.append(("TRI annualisé", _pct(ar["tri_annualise"])))
    etapes = [
        ("Produit net de vente", ar["produit_net_vente"], True),
        ("Coût d'acquisition", -ar["cout_total_acquisition"], False),
        ("Frais de portage", -ar["frais_portage_total"], False),
        ("Impôt", -ar["impot_total"], False),
        ("Marge nette", ar["marge_nette"], True),
    ]
    image = charts.chart_pont_marge(etapes)
    _ajouter_table_et_graphique(doc, lignes, image, largeur_table_cm=13.0, largeur_image_cm=12.3)


def _section_endettement(doc, inp, payload, resultat, is_achat_revente):
    mensualite_projet, loyers_mensuels = _mensualite_et_loyers(inp, resultat, is_achat_revente)
    r = endet_mod.calculer_taux_endettement(
        payload.profil.revenus_nets_mensuels_foyer,
        payload.profil.autres_revenus_mensuels,
        payload.profil.mensualites_credits_existants,
        mensualite_projet,
        loyers_mensuels,
    )
    lignes = [
        ("Mensualité du projet retenue (hors assurance)", _eur(mensualite_projet)),
        ("Recettes locatives prévisionnelles retenues à 70 %", _eur(loyers_mensuels * endet_mod.PONDERATION_LOYERS)),
        ("Revenus considérés par la banque", _eur(r.revenus_consideres_mensuels)),
        ("Mensualités totales (crédits existants + projet)", _eur(r.mensualites_totales_mensuelles)),
        ("Taux d'endettement", _pct(r.taux_endettement, 1)),
        ("Seuil HCSF", _pct(r.seuil_hcsf, 0)),
    ]
    image = charts.chart_barres(
        [("Taux d'endettement du projet", r.taux_endettement)],
        "Taux d'endettement vs seuil HCSF",
        formatter=lambda v: _pct(v, 1),
        seuil=r.seuil_hcsf,
        seuil_label=f"Seuil {_pct(r.seuil_hcsf, 0)}",
    )
    _ajouter_table_et_graphique(doc, lignes, image)
    p = doc.add_paragraph()
    run = p.add_run(
        "⚠ Dépasse le seuil HCSF de 35 %." if r.depasse_seuil else "OK — sous le seuil HCSF de 35 %."
    )
    run.bold = True
    run.font.color.rgb = ALERTE_COLOR if r.depasse_seuil else PRIMARY_COLOR
    if not r.depasse_seuil:
        marge = doc.add_paragraph(
            f"Marge avant d'atteindre le seuil : {_eur(r.marge_avant_seuil)}/mois de mensualité supplémentaire supportable."
        )
        marge.runs[0].font.size = Pt(9)


def _section_avertissements(doc, avertissements):
    for texte in avertissements:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(texte)
        run.font.color.rgb = RGBColor(0x99, 0x5C, 0x00)


def _section_mentions(doc):
    doc.add_paragraph(
        "Ce document est une estimation pédagogique générée automatiquement à partir des "
        "hypothèses saisies par l'utilisateur (prix, loyers ou tarifs de location, charges, taux, durée...). Il ne "
        "constitue ni une offre de prêt, ni un conseil fiscal ou juridique personnalisé."
    )
    doc.add_paragraph(
        "Le barème de l'impôt sur le revenu, les taux de prélèvements sociaux et les règles "
        "d'amortissement appliqués correspondent à la législation en vigueur au moment de la "
        "génération du document et sont susceptibles d'évoluer."
    )
    note = doc.add_paragraph(
        "Ces éléments doivent être vérifiés et validés par un professionnel (notaire, "
        "expert-comptable, courtier en crédit) avant toute décision d'investissement ou "
        "constitution d'un dossier de financement définitif."
    )
    note.runs[0].italic = True


# ---------------------------------------------------------------------------
# Point d'entrée
# ---------------------------------------------------------------------------

def generer_dossier_word(payload: ExportDossierInput) -> bytes:
    inp = payload.simulation
    resultat = clean_result(simuler(inp))
    is_achat_revente = inp.type_projet == TypeProjet.achat_revente
    is_lcd = inp.type_projet == TypeProjet.location_courte_duree
    is_meublee = is_lcd or inp.regime_location.value == "meublee"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = POLICE
    style.font.size = Pt(10.5)

    _configurer_page(doc)
    _configurer_entete_pied(doc)
    _ajouter_page_de_garde(doc, payload, inp)

    annee1 = None if is_achat_revente else resultat["annees"][0]
    meilleur_regime = None if is_achat_revente else _meilleur_regime(annee1)

    sections: list[tuple[str, "callable"]] = [
        (
            "Synthèse du projet",
            lambda d: _section_synthese(d, payload, inp, resultat, is_achat_revente, meilleur_regime, annee1),
        ),
        (
            "Présentation du bien et du projet",
            lambda d: _section_presentation(d, payload, inp, resultat, is_achat_revente),
        ),
    ]
    if payload.profil is not None:
        sections.append(("Profil de l'emprunteur", lambda d: _section_profil(d, payload)))
    sections.append(("Plan de financement", lambda d: _section_financement(d, inp, resultat, is_achat_revente)))
    if is_achat_revente:
        sections.append(("Détail de l'opération d'achat-revente", lambda d: _section_achat_revente_detail(d, inp, resultat)))
    else:
        sections.append(
            ("Recettes et charges d'exploitation annuelles", lambda d: _section_charges(d, inp, annee1, is_meublee, is_lcd))
        )
    if payload.profil is not None:
        sections.append(
            ("Taux d'endettement (HCSF)", lambda d: _section_endettement(d, inp, payload, resultat, is_achat_revente))
        )
    avertissements = resultat.get("avertissements") or []
    if avertissements:
        sections.append(("Points d'attention", lambda d: _section_avertissements(d, avertissements)))
    sections.append(("Mentions et méthodologie", lambda d: _section_mentions(d)))

    _nouvelle_page_chapitre(doc)
    _ajouter_sommaire(doc, [titre for titre, _ in sections])

    for i, (titre, fn) in enumerate(sections, start=1):
        _nouvelle_page_chapitre(doc)
        _ajouter_titre(doc, f"{i}. {titre}")
        fn(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
