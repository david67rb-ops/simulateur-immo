"""Repère de prix saisi par l'utilisateur dans les zones sans DVF
(Alsace-Moselle, Mayotte). Lancement :

    venv/bin/python -m unittest discover -s tests
"""
from __future__ import annotations

import io
import json
import sys
import unittest
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402

from app import analyse, dossier_export, projet_signe, schemas  # noqa: E402
from app.chapitres_dossier import FORMULES_DOSSIER  # noqa: E402
from app.simulation import simuler  # noqa: E402
from app.utils import clean_result  # noqa: E402
from gui.state import default_market_state  # noqa: E402

AUJOURD_HUI = date(2026, 10, 15)


def _etat(**valeurs) -> dict:
    etat = default_market_state()
    etat.update({"prix_ref_date": AUJOURD_HUI.isoformat(), **valeurs})
    return etat


def _inp(prix: float = 150_000, surface: float = 45) -> schemas.SimulationInput:
    return schemas.SimulationInput(
        surface_m2=surface, prix_achat=prix, frais_notaire=prix * 0.08, loyer_mensuel_hors_charges=700, apport=15_000
    )


class EcartAuRepere(unittest.TestCase):
    def test_au_dessus_et_en_dessous(self):
        self.assertAlmostEqual(analyse.ecart_au_repere(2_750, 2_500), 0.10)
        self.assertAlmostEqual(analyse.ecart_au_repere(2_450, 2_600), -0.0576923, places=6)

    def test_meme_prix_au_m2_que_pour_dvf(self):
        self.assertAlmostEqual(analyse.prix_projet_m2(150_000, 45), 3_333.333, places=3)
        self.assertIsNone(analyse.prix_projet_m2(150_000, 0))


class SeuilsAdresse(unittest.TestCase):
    def test_vert_orange_rouge(self):
        self.assertEqual(analyse.niveau_ecart_repere(-0.20, "adresse"), "vert")
        self.assertEqual(analyse.niveau_ecart_repere(0.029, "adresse"), "vert")
        self.assertEqual(analyse.niveau_ecart_repere(analyse.SEUIL_PRIX_VERT, "adresse"), "orange")
        self.assertEqual(analyse.niveau_ecart_repere(0.099, "adresse"), "orange")
        self.assertEqual(analyse.niveau_ecart_repere(analyse.SEUIL_PRIX_ORANGE, "adresse"), "rouge")

    def test_memes_seuils_que_dvf(self):
        for ecart in (-0.1, 0.0, 0.05, 0.2):
            self.assertEqual(analyse.niveau_ecart_repere(ecart, "adresse"), analyse.niveau_ecart_au_marche(ecart))


class NiveauCommune(unittest.TestCase):
    def test_jamais_rouge(self):
        self.assertEqual(analyse.niveau_ecart_repere(0.40, "commune"), "orange")
        self.assertEqual(analyse.niveau_ecart_repere(0.05, "commune"), "orange")
        self.assertEqual(analyse.niveau_ecart_repere(-0.10, "commune"), "vert")

    def test_verdict_affiche_le_message_communal(self):
        inp = _inp(prix=3_000 * 45)
        resultat = clean_result(simuler(inp))
        v = analyse.verdict_global(inp, resultat, 0.25, niveau_repere="commune")
        prix = next(c for c in v["criteres"] if c["nom"] == "Prix")
        self.assertEqual(prix["niveau"], "orange")
        self.assertIn(analyse.MESSAGE_REPERE_COMMUNAL, prix["texte"])


class Validation(unittest.TestCase):
    def valider(self, moyen, bas=None, haut=None, jour=AUJOURD_HUI):
        return analyse.valider_prix_reference(moyen, bas, haut, jour.isoformat(), AUJOURD_HUI)

    def test_bornes(self):
        self.assertEqual(self.valider(300), ([], []))
        self.assertEqual(self.valider(20_000), ([], []))
        self.assertEqual(self.valider(150_000)[0], ["Ce chiffre ressemble à un prix total, pas à un prix au m²."])
        self.assertEqual(self.valider(299)[0], ["Ce prix au m² paraît trop bas : vérifie la saisie."])

    def test_ordre_de_la_fourchette(self):
        self.assertEqual(self.valider(2_600, 2_200, 3_000), ([], []))
        self.assertTrue(self.valider(2_600, 2_700, 3_000)[0])
        self.assertTrue(self.valider(2_600, 2_200, 2_500)[0])

    def test_date_dans_le_futur(self):
        erreurs, _ = self.valider(2_600, jour=AUJOURD_HUI + timedelta(days=1))
        self.assertEqual(erreurs, ["La date du relevé ne peut pas être dans le futur."])

    def test_repere_invalide_ignore(self):
        self.assertIsNone(analyse.repere_prix(_etat(prix_ref_m2=150_000), AUJOURD_HUI))


class AncienneteDuReleve(unittest.TestCase):
    def test_plus_de_6_mois(self):
        erreurs, avertissements = analyse.valider_prix_reference(
            2_600, date_releve=(AUJOURD_HUI - timedelta(days=200)).isoformat(), aujourd_hui=AUJOURD_HUI
        )
        self.assertEqual(erreurs, [])
        self.assertEqual(
            avertissements, ["Relevé de plus de 6 mois : mets-le à jour avant d'envoyer le dossier à la banque."]
        )
        # L'avertissement n'empêche pas d'utiliser le repère.
        self.assertIsNotNone(
            analyse.repere_prix(_etat(prix_ref_m2=2_600, prix_ref_date="2026-03-01"), AUJOURD_HUI)
        )

    def test_moins_de_6_mois(self):
        _, avertissements = analyse.valider_prix_reference(2_600, date_releve="2026-06-01", aujourd_hui=AUJOURD_HUI)
        self.assertEqual(avertissements, [])


class SansRepere(unittest.TestCase):
    def test_critere_non_evalue(self):
        self.assertIsNone(analyse.repere_prix(_etat(), AUJOURD_HUI))
        inp = _inp()
        v = analyse.verdict_global(inp, clean_result(simuler(inp)), None, raison_prix=analyse.RAISON_SANS_REPERE)
        prix = next(c for c in v["criteres"] if c["nom"] == "Prix")
        self.assertEqual(prix, {"nom": "Prix", "niveau": "neutre", "texte": analyse.RAISON_SANS_REPERE})


class AncienFichierProjet(unittest.TestCase):
    def test_chargement_sans_les_nouveaux_champs(self):
        ancien = {k: v for k, v in default_market_state().items() if not k.startswith("prix_ref_")}
        ancien["adresse"] = "1 place Kléber, 67000 Strasbourg"
        fichier = projet_signe.exporter({"marche": ancien}, "dossier-test", None, AUJOURD_HUI)
        self.assertNotIn("prix_ref_m2", json.loads(fichier)["etat"]["marche"])

        contenu = projet_signe.importer(fichier)  # signature toujours valide
        etat = default_market_state()
        projet_signe.reprendre(etat, contenu["etat"]["marche"])
        self.assertEqual(etat["adresse"], "1 place Kléber, 67000 Strasbourg")
        self.assertIsNone(etat["prix_ref_m2"])
        self.assertEqual(etat["prix_ref_niveau"], "adresse")
        self.assertEqual(etat["prix_ref_source"], "meilleursagents")

    def test_le_repere_ne_change_pas_l_identifiant_du_dossier(self):
        # Le dossier ouvert (30 jours) suit son identifiant, pas le contenu du projet.
        etat = _etat(prix_ref_m2=3_400)
        a = json.loads(projet_signe.exporter({"marche": etat}, "dossier-x", "2026-11-14", AUJOURD_HUI))
        etat["prix_ref_m2"] = 3_100
        b = json.loads(projet_signe.exporter({"marche": etat}, "dossier-x", "2026-11-14", AUJOURD_HUI))
        self.assertEqual(a["dossier_id"], b["dossier_id"])


MOTS_DE_VERDICT = ("vert", "orange", "rouge", "à revoir", "bon prix", "favorable", "à surveiller", "solide", "Critère")


def _marche(repere: dict | None) -> dict:
    return {
        "adresse": "1 place Kléber, 67000 Strasbourg",
        "lat": 48.5834,
        "lon": 7.7458,
        "code_insee": "67482",
        "code_departement": "67",
        "type_bien": "appartement",
        "comparables": {"nb_transactions": 0, "prix_m2_moyen": None, "indisponible": "…"},
        "sans_dvf": True,
        "loyer": None,
        "commune_indicateurs": None,
        "repere_prix": repere,
    }


def _texte_docx(contenu: bytes) -> str:
    doc = Document(io.BytesIO(contenu))
    morceaux = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for ligne in table.rows:
            morceaux.extend(c.text for c in ligne.cells)
    return "\n".join(morceaux)


class TexteDuDossier(unittest.TestCase):
    repere = analyse.repere_prix(
        _etat(prix_ref_m2=2_600, prix_ref_bas_m2=2_200, prix_ref_haut_m2=3_000, prix_ref_date="2026-10-15"),
        AUJOURD_HUI,
    )
    inp = _inp(prix=2_450 * 40, surface=40)

    def test_texte_banque_factuel(self):
        textes = dossier_export.texte_repere_prix(self.inp, _marche(self.repere), pour_banque=True)
        self.assertEqual(
            textes[0],
            "Repère de prix : 2 600 €/m² (fourchette de 2 200 à 3 000 €/m²), estimation "
            "MeilleursAgents à l'adresse du bien, relevée par l'emprunteur le 15/10/2026. Prix du projet : "
            "2 450 €/m², soit 5,8 % sous ce repère.",
        )
        self.assertIn("régime du Livre foncier", textes[1])
        self.assertIn("à rapprocher d'un avis de valeur professionnel", textes[1])
        for mot in MOTS_DE_VERDICT:
            self.assertNotIn(mot, " ".join(textes))

    def test_sans_fourchette(self):
        repere = {**self.repere, "bas_m2": None, "haut_m2": None}
        texte = dossier_export.texte_repere_prix(self.inp, _marche(repere), pour_banque=True)[0]
        self.assertNotIn("fourchette", texte)

    def test_au_dessus_du_repere_communal(self):
        repere = {**self.repere, "niveau": "commune", "source": "agence"}
        inp = _inp(prix=2_860 * 40, surface=40)
        texte = dossier_export.texte_repere_prix(inp, _marche(repere), pour_banque=True)[0]
        self.assertIn("estimation d'agence moyenne de la commune", texte)
        self.assertIn("soit 10,0 % au-dessus de ce repère", texte)

    def test_sans_repere(self):
        textes = dossier_export.texte_repere_prix(self.inp, _marche(None), pour_banque=True)
        self.assertEqual(len(textes), 1)
        self.assertTrue(textes[0].endswith("Aucun repère de prix n'a été renseigné."))

    def test_les_trois_formules(self):
        resultat = clean_result(simuler(self.inp))
        for cle, (_nom, chapitres) in FORMULES_DOSSIER.items():
            with self.subTest(formule=cle):
                payload = schemas.ExportDossierInput(
                    simulation=self.inp, marche=_marche(self.repere), chapitres=sorted(chapitres)
                )
                texte = _texte_docx(dossier_export.generer_dossier_word(payload))
                self.assertIn("Repère de prix : 2 600 €/m²", texte)
                self.assertIn("relevée par l'emprunteur le 15/10/2026", texte)
                personnel = cle == "personnel"
                self.assertEqual("Critère Prix sur Credaura" in texte, personnel)
                if not personnel:
                    chapitre = dossier_export.texte_repere_prix(self.inp, _marche(self.repere), pour_banque=True)
                    for mot in MOTS_DE_VERDICT:
                        self.assertNotIn(mot, " ".join(chapitre))
        self.assertIsNotNone(resultat)


if __name__ == "__main__":
    unittest.main()
