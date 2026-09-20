"""Génération du dossier de financement au format Word (.docx)."""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from . import endettement as endet_mod
from .schemas import ExportDossierInput, TypeProjet
from .simulation import simuler
from .utils import clean_result

PRIMARY_COLOR = RGBColor(0x1D, 0x6F, 0x5C)


def _eur(v: float) -> str:
    return f"{v:,.0f} €".replace(",", " ")


def _pct(v: float, digits: int = 1) -> str:
    return f"{v * 100:.{digits}f} %"


def _ajouter_titre(doc: Document, texte: str, niveau: int = 1):
    h = doc.add_heading(texte, level=niveau)
    for run in h.runs:
        run.font.color.rgb = PRIMARY_COLOR


def _ajouter_table_kv(doc: Document, lignes: list[tuple[str, str]]):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for cle, valeur in lignes:
        row = table.add_row()
        row.cells[0].text = cle
        row.cells[1].text = valeur
    return table


def generer_dossier_word(payload: ExportDossierInput) -> bytes:
    inp = payload.simulation
    resultat = clean_result(simuler(inp))

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    titre = doc.add_heading("Dossier de financement immobilier", level=0)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sous_titre = doc.add_paragraph(
        "Document généré automatiquement — estimation pédagogique, à faire valider "
        "par un professionnel avant toute décision."
    )
    sous_titre.runs[0].italic = True

    # --- 1. Présentation du projet ---
    _ajouter_titre(doc, "1. Présentation du projet")
    _ajouter_table_kv(
        doc,
        [
            ("Type de projet", inp.type_projet.value.replace("_", " ")),
            ("Structure juridique", inp.structure_juridique.value.replace("_", " ")),
            ("Type de bien", inp.type_bien.value),
            ("Surface", f"{inp.surface_m2:.0f} m²"),
            ("Prix d'achat", _eur(inp.prix_achat)),
            ("Frais de notaire", _eur(inp.frais_notaire)),
            ("Montant travaux", _eur(inp.montant_travaux)),
        ],
    )

    # --- 2. Plan de financement ---
    _ajouter_titre(doc, "2. Plan de financement")
    lignes_financement = [
        ("Coût total d'acquisition", _eur(resultat.get("cout_total_acquisition", 0))),
        ("Apport personnel", _eur(resultat.get("apport_reel", 0))),
        ("Montant emprunté", _eur(resultat.get("montant_emprunte", 0))),
        ("Taux du crédit", _pct(inp.taux_credit_annuel, 2)),
        ("Durée du crédit", f"{inp.duree_credit_annees} ans"),
    ]
    if inp.differe_type.value != "aucun":
        lignes_financement.append(
            ("Différé de crédit", f"{inp.differe_type.value} — {inp.differe_duree_mois} mois")
        )
        lignes_financement.append(
            (
                "Mensualité 1ère année",
                _eur(resultat.get("mensualite_annee1_hors_assurance", 0)) + "/mois",
            )
        )
    lignes_financement.append(
        ("Mensualité en régime de croisière", _eur(resultat.get("mensualite_credit_hors_assurance", 0)) + "/mois")
    )
    _ajouter_table_kv(doc, lignes_financement)

    # --- 3. Rentabilité prévisionnelle ---
    if inp.type_projet != TypeProjet.achat_revente:
        _ajouter_titre(doc, "3. Rentabilité prévisionnelle")
        _ajouter_table_kv(
            doc,
            [
                ("Rendement brut", _pct(resultat.get("rendement_brut", 0))),
                ("Rendement net de charges", _pct(resultat.get("rendement_net_charges", 0))),
            ],
        )
        annee1 = resultat["annees"][0]
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = (
            "Régime",
            "Revenu imposable",
            "Impôt total",
            "Cash-flow mensuel net",
        )
        for regime, fiscal in annee1["fiscal"].items():
            row = table.add_row().cells
            row[0].text = regime
            row[1].text = _eur(fiscal["revenu_imposable"])
            row[2].text = _eur(fiscal["total_prelevements"])
            row[3].text = _eur(annee1["cashflow_apres_impot"][regime] / 12)
    else:
        _ajouter_titre(doc, "3. Rentabilité de l'opération (achat-revente)")
        ar = resultat["achat_revente"]
        _ajouter_table_kv(
            doc,
            [
                ("Marge brute avant impôt", _eur(ar["marge_brute_avant_impot"])),
                ("Régime fiscal", ar["regime_fiscal"]),
                ("Impôt total", _eur(ar["impot_total"])),
                ("Marge nette", _eur(ar["marge_nette"])),
                ("Rentabilité de l'opération", _pct(ar["rentabilite_operation_pct"])),
            ],
        )

    # --- 4. Profil emprunteur & taux d'endettement ---
    if payload.profil is not None:
        _ajouter_titre(doc, "4. Profil emprunteur et taux d'endettement")
        if inp.type_projet == TypeProjet.achat_revente:
            ar = resultat["achat_revente"]
            mensualite_projet = ar["frais_portage_interets"] / inp.duree_portage_mois
            loyers_mensuels = 0.0
        else:
            mensualite_projet = resultat.get("mensualite_credit_hors_assurance", 0)
            loyers_mensuels = resultat["annees"][0]["loyers_bruts"] / 12

        r = endet_mod.calculer_taux_endettement(
            payload.profil.revenus_nets_mensuels_foyer,
            payload.profil.autres_revenus_mensuels,
            payload.profil.mensualites_credits_existants,
            mensualite_projet,
            loyers_mensuels,
        )
        _ajouter_table_kv(
            doc,
            [
                ("Revenus nets mensuels du foyer", _eur(payload.profil.revenus_nets_mensuels_foyer)),
                ("Autres revenus mensuels", _eur(payload.profil.autres_revenus_mensuels)),
                ("Mensualités de crédits existants", _eur(payload.profil.mensualites_credits_existants)),
                ("Mensualité du projet (hors assurance)", _eur(mensualite_projet)),
                ("Loyers prévisionnels retenus (70 %)", _eur(loyers_mensuels * endet_mod.PONDERATION_LOYERS)),
                ("Revenus considérés par la banque", _eur(r.revenus_consideres_mensuels)),
                ("Taux d'endettement", _pct(r.taux_endettement)),
                ("Seuil HCSF", _pct(r.seuil_hcsf)),
                (
                    "Statut",
                    "⚠ Dépasse le seuil HCSF" if r.depasse_seuil else "OK, sous le seuil HCSF",
                ),
            ],
        )

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Ce document est une estimation pédagogique générée automatiquement. Les "
        "hypothèses fiscales et bancaires doivent être vérifiées par un professionnel "
        "(notaire, comptable, courtier) avant toute décision d'investissement."
    )
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(9)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
